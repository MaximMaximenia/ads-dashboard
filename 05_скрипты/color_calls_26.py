import openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

src='/Users/maksim/Desktop/Реклама/03_анализ/Звонки_ДАСТО_26.10.2026.xlsx'
wb=openpyxl.load_workbook(src)
ws=wb['Звонки 26.10']
# insert columns G:I
ws.insert_cols(7,3)
headers=['Цена, BYN','Выручка (70%), BYN','Потенциал (70%), BYN']
for i,h in enumerate(headers,7):
    c=ws.cell(1,i,h)
    c.fill=PatternFill('solid',fgColor='17181A'); c.font=Font(color='FFFFFF',bold=True,size=10)
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
# status colors
fills={
 'Приедет':('C6EFCE','006100'),
 'Под вопросом':('FFEB9C','9C6500'),
 'Недозвон':('D9D9D9','404040'),
 'Не по теме':('FFC7CE','9C0006'),
 'Старый / не по теме':('FCE4D6','833C00'),
}
price_for=lambda topic: 80 if 'Развал' in topic else (30 if 'Регулировка' in topic else None)
for r in range(2,ws.max_row+1):
    status=ws.cell(r,5).value; topic=ws.cell(r,4).value
    # clear existing fills row
    for c in range(1,10):
        ws.cell(r,c).fill=PatternFill(fill_type=None)
    bg,fg=fills.get(status,('FFFFFF','000000'))
    for c in range(1,10):
        cell=ws.cell(r,c); cell.fill=PatternFill('solid',fgColor=bg); cell.font=Font(color=fg,size=9)
        cell.alignment=Alignment(vertical='center',wrap_text=True)
        cell.border=Border(bottom=Side(style='thin',color='FFFFFF'))
    price=price_for(topic)
    ws.cell(r,7).value=price
    ws.cell(r,7).number_format='#,##0'
    if status=='Приедет':
        ws.cell(r,8).value=f'=IF(AND(G{r}<>"",E{r}="Приедет"),G{r}*0.7,0)'
        ws.cell(r,9).value=0
    elif status=='Под вопросом':
        ws.cell(r,8).value=0
        ws.cell(r,9).value=f'=IF(G{r}<>"",G{r}*0.7,0)'
    else:
        ws.cell(r,8).value=0; ws.cell(r,9).value=0
    for c in (8,9): ws.cell(r,c).number_format='#,##0'
# totals row
last=ws.max_row
tot=last+1
ws.cell(tot,1,'ИТОГО').font=Font(bold=True,size=10)
ws.merge_cells(start_row=tot,start_column=1,end_row=tot,end_column=6)
ws.cell(tot,1).alignment=Alignment(horizontal='right',vertical='center')
ws.cell(tot,7,'=SUM(G2:G%d)'%last).font=Font(bold=True,size=10)
ws.cell(tot,8,'=SUM(H2:H%d)'%last).font=Font(bold=True,size=10)
ws.cell(tot,9,'=SUM(I2:I%d)'%last).font=Font(bold=True,size=10)
for c in (7,8,9): ws.cell(tot,c).number_format='#,##0'
for c in range(1,10):
    ws.cell(tot,c).fill=PatternFill('solid',fgColor='17181A'); ws.cell(tot,c).font=Font(color='FFFFFF',bold=True,size=10)
# widths
for col,w in zip('ABCDEFGHI',[8,7,15,27,21,50,11,15,16]): ws.column_dimensions[col].width=w
ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:I{last}'

# summary
s=wb['Сводка']
# clear and rebuild
wb.remove(s); s=wb.create_sheet('Сводка')
s['A1']='ДАСТО — звонки 26.10.2026'; s['A1'].font=Font(size=18,bold=True)
s['A2']='Входящие из A1 ВАТС. Статусы и расшифровка — предварительные.'; s['A2'].font=Font(size=10,color='6E6E73')
r=4
s.cell(r,1,'Показатель').font=Font(bold=True,color='FFFFFF'); s.cell(r,1).fill=PatternFill('solid',fgColor='17181A')
s.cell(r,2,'Значение').font=Font(bold=True,color='FFFFFF'); s.cell(r,2).fill=PatternFill('solid',fgColor='17181A')
r+=1
def add(k,v,fmt=None,bold=False):
    global r
    s.cell(r,1,k).font=Font(bold=bold)
    s.cell(r,2,v).font=Font(bold=bold)
    if fmt: s.cell(r,2).number_format=fmt
    r+=1
add('Всего входящих звонков',f"=COUNTA('Звонки 26.10'!A2:A{last})")
add('Уникальных номеров',42)
add('Приедет',f"=COUNTIF('Звонки 26.10'!E2:E{last},\"Приедет\")",bold=True)
add('Под вопросом',f"=COUNTIF('Звонки 26.10'!E2:E{last},\"Под вопросом\")")
add('Недозвон',f"=COUNTIF('Звонки 26.10'!E2:E{last},\"Недозвон\")")
add('Не по теме / старые',f"=COUNTIF('Звонки 26.10'!E2:E{last},\"Не по теме\")+COUNTIF('Звонки 26.10'!E2:E{last},\"Старый / не по теме\")")
r+=1
add('Ожидаемая выручка по известным ценам, BYN',f"=SUM('Звонки 26.10'!H2:H{last})",'#,##0',True)
add('Потенциал по «под вопросом» (известные цены), BYN',f"=SUM('Звонки 26.10'!I2:I{last})",'#,##0')
add('Сумма цен по известным услугам, BYN',f"=SUM('Звонки 26.10'!G2:G{last})",'#,##0')
r+=1
s.cell(r,1,'70% — доля доезжаемости. Выручка = цена × 0,7 только для «Приедет».').font=Font(size=9,color='6E6E73'); r+=1
s.cell(r,1,'Развал: 80 BYN (60–100). Регулировка фар: 30 BYN. Остальные цены заполняет владелец.').font=Font(size=9,color='6E6E73'); r+=1
s.cell(r,1,'После заполнения колонки G «Цена» колонки H и I пересчитаются автоматически.').font=Font(size=9,color='6E6E73')
s.column_dimensions['A'].width=58; s.column_dimensions['B'].width=16
for rr in range(4, r+1):
    s.cell(rr,1).alignment=Alignment(vertical='center',wrap_text=True)
    s.cell(rr,2).alignment=Alignment(vertical='center')

# README
rd=wb.create_sheet('Как заполнять')
rd['A1']='Как заполнять таблицу'; rd['A1'].font=Font(size=16,bold=True)
lines=[
 '1. Колонка G «Цена, BYN» — цена услуги. Развал 80 (60–100), регулировка фар 30.',
 '2. Колонки H/I пересчитаются: H — ожидаемая выручка по «Приедет» (70%), I — потенциал по «Под вопросом».',
 '3. Для остальных услуг (масло, подвеска, тормоза, глушитель и т.д.) впишите цену в G.',
 '4. Статусы: зелёный — приедет, жёлтый — под вопросом, серый — недозвон, красный — не по теме, оранжевый — старый клиент.',
 '5. Итог — внизу листа «Звонки 26.10» и на «Сводке».',
]
for i,l in enumerate(lines,3):
    rd.cell(i,1,l).font=Font(size=11); rd.cell(i,1).alignment=Alignment(wrap_text=True,vertical='top')
rd.column_dimensions['A'].width=110
wb.save(src)
print('saved',src,'rows',last,'total',tot)
