from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG=RGBColor(0xF7,0xF7,0xF5); DARK=RGBColor(0x17,0x18,0x1A); RED=RGBColor(0xD0,0x13,0x0D)
GRAY=RGBColor(0x6E,0x6E,0x73); LINE=RGBColor(0xE3,0xE3,0xDF); WHITE=RGBColor(0xFF,0xFF,0xFF); SOFT=RGBColor(0xEF,0xEF,0xEC)
GREEN=RGBColor(0x1F,0x8A,0x4C); FONT="Arial"
p=Presentation(); p.slide_width=Inches(13.333); p.slide_height=Inches(7.5)

def slide(kicker,title,footer="Регулировка фар · 26–27.10.2026 · без изменений в аккаунте"):
 s=p.slides.add_slide(p.slide_layouts[6]); bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,p.slide_width,p.slide_height)
 bg.fill.solid(); bg.fill.fore_color.rgb=BG; bg.line.fill.background()
 bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(0.10),p.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb=RED; bar.line.fill.background()
 t=s.shapes.add_textbox(Inches(0.55),Inches(0.30),Inches(12),Inches(0.25)); r=t.text_frame.paragraphs[0].add_run(); r.text=kicker; r.font.name=FONT; r.font.size=Pt(9); r.font.bold=True; r.font.color.rgb=RED
 t=s.shapes.add_textbox(Inches(0.55),Inches(0.56),Inches(12.2),Inches(0.6)); r=t.text_frame.paragraphs[0].add_run(); r.text=title; r.font.name=FONT; r.font.size=Pt(23); r.font.bold=True; r.font.color.rgb=DARK
 t=s.shapes.add_textbox(Inches(0.55),Inches(7.12),Inches(12.2),Inches(0.2)); r=t.text_frame.paragraphs[0].add_run(); r.text=footer; r.font.name=FONT; r.font.size=Pt(8); r.font.color.rgb=GRAY
 return s

def text(s,x,y,w,h,txt,size=11,bold=False,color=DARK):
 box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.word_wrap=True
 tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
 for i,line in enumerate(txt if isinstance(txt,list) else [txt]):
  para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.space_after=Pt(4)
  r=para.add_run(); r.text=line; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color

def bullets(s,x,y,w,h,items,size=10,gap=7):
 box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.word_wrap=True
 tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
 for i,item in enumerate(items):
  para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.space_after=Pt(gap); para.line_spacing=1.02
  r0=para.add_run(); r0.text="●  "; r0.font.name=FONT; r0.font.size=Pt(size-3); r0.font.color.rgb=RED
  r=para.add_run(); r.text=item; r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=DARK

def kpi(s,x,y,w,val,label,dark=False):
 card=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(1.0)); card.adjustments[0]=0.10
 card.fill.solid(); card.fill.fore_color.rgb=DARK if dark else WHITE; card.line.color.rgb=DARK if dark else LINE
 t=s.shapes.add_textbox(Inches(x+0.16),Inches(y+0.10),Inches(w-0.32),Inches(0.44)); r=t.text_frame.paragraphs[0].add_run(); r.text=val
 r.font.name=FONT; r.font.size=Pt(25); r.font.bold=True; r.font.color.rgb=WHITE if dark else RED
 t=s.shapes.add_textbox(Inches(x+0.16),Inches(y+0.62),Inches(w-0.32),Inches(0.25)); r=t.text_frame.paragraphs[0].add_run(); r.text=label
 r.font.name=FONT; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0xD9,0xD9,0xD9) if dark else GRAY

def table(s,x,y,w,h,data,widths,font=8.6,header_font=8.0,row_h=0.34):
 rows,cols=len(data),len(data[0]); shp=s.shapes.add_table(rows,cols,Inches(x),Inches(y),Inches(w),Inches(h)); tbl=shp.table
 for i,cw in enumerate(widths): tbl.columns[i].width=Inches(cw)
 for r in range(rows):
  tbl.rows[r].height=Inches(row_h)
  for c in range(cols):
   cell=tbl.cell(r,c); cell.margin_left=Inches(0.05); cell.margin_right=Inches(0.05); cell.margin_top=Inches(0.01); cell.margin_bottom=Inches(0.01)
   cell.vertical_anchor=MSO_ANCHOR.MIDDLE; cell.fill.solid()
   cell.fill.fore_color.rgb=DARK if r==0 else (WHITE if r%2 else SOFT)
   para=cell.text_frame.paragraphs[0]; para.alignment=PP_ALIGN.RIGHT if c>0 else PP_ALIGN.LEFT
   run=para.add_run(); run.text=str(data[r][c]); run.font.name=FONT; run.font.size=Pt(header_font if r==0 else font); run.font.bold=(r==0)
   run.font.color.rgb=WHITE if r==0 else DARK

def title_slide():
 s=p.slides.add_slide(p.slide_layouts[6]); bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,p.slide_width,p.slide_height)
 bg.fill.solid(); bg.fill.fore_color.rgb=DARK; bg.line.fill.background()
 bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(0.70),Inches(1.25),Inches(0.12),Inches(4.0)); bar.fill.solid(); bar.fill.fore_color.rgb=RED; bar.line.fill.background()
 text(s,1.05,1.20,10.8,0.3,"GOOGLE ADS / ДАСТО · РЕГУЛИРОВКА ФАР",11,True,RED)
 text(s,1.05,1.65,10.8,1.0,"Регулировка фар\n26–27 октября",42,True,WHITE)
 text(s,1.05,3.35,10.8,0.75,"Сколько трафика выкупили, почём заявка и где потери",18,False,RGBColor(0xE6,0xE6,0xE6))
 line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1.05),Inches(5.15),Inches(5.0),Inches(0.02)); line.fill.solid(); line.fill.fore_color.rgb=RGBColor(0x3A,0x3A,0x3C); line.line.fill.background()
 text(s,1.05,5.33,10.8,0.3,"Подготовлено 28.10.2026 · без изменений в аккаунте",10,False,RGBColor(0xB9,0xB9,0xB9))

# 1 title
title_slide()
# 2 summary
s=slide("КОРОТКО","10 заявок за $17,28: CPA $1,73")
kpi(s,0.55,1.35,2.8,"10","заявок за 2 дня",False)
kpi(s,3.50,1.35,2.8,"$1,73","CPA за 2 дня",True)
kpi(s,6.45,1.35,2.8,"20,8%","конверсия клика")
kpi(s,9.40,1.35,3.0,"$16","текущая целевая CPA")
bullets(s,0.55,2.72,6.0,3.9,[
 "26.10: 6 заявок, $9,02, CPA $1,50.",
 "27.10: 4 заявки, $8,26, CPA $2,07.",
 "48 кликов, 297 показов, CPC $0,36; 98% кликов — мобильные.",
 "Телефонные звонки: 6 и 3 подтверждённых; стоимость подтверждённого звонка $1,50 / $2,75.",
],size=9.6,gap=8)
text(s,6.85,2.72,5.9,0.3,"ЧТО ЭТО ЗНАЧИТ",9.2,True,RED)
bullets(s,6.85,3.05,5.9,3.5,[
 "Трафик дешёвый: клик $0,36, конверсия 21%.",
 "Но 40% показов потеряно из-за ранга; бюджет не режет — lostB 0%.",
 "Это не история про «мало денег», это история про видимость.",
 "Снижать целевую CPA до $2–3 опасно: можно урезать объём, не улучшив ранг.",
],size=9.4,gap=7)

# 3 daily table
s=slide("ДНИ","26 и 27 октября: цифры по дням")
d=[["День","Показы","Клики","Расход","Заявки","CTR","CPC","CVR","CPA"],
   ["26.10","148","25","$9,02","6","16,9%","$0,36","24,0%","$1,50"],
   ["27.10","149","23","$8,26","4","15,4%","$0,36","17,4%","$2,07"],
   ["Итого","297","48","$17,28","10","16,2%","$0,36","20,8%","$1,73"]]
table(s,0.55,1.35,12.2,2.0,d,[2.0,1.2,1.0,1.5,1.2,1.2,1.2,1.2,1.7],font=9.2,header_font=8.2,row_h=0.38)
bullets(s,0.55,3.55,12.2,2.8,[
 "26-е: 6 заявок, CPA $1,50 — лучший день.",
 "27-е: заявок меньше (4), CPA $2,07 — всё ещё дёшево.",
 "Расход почти ровный: $9,02 vs $8,26.",
 "Клики: 25 vs 23; показы: 148 vs 149.",
 "Конверсия просела: 24% → 17,4%, но выборка маленькая.",
],size=9.6,gap=8)

# 4 traffic
s=slide("ТРАФИК","Что именно выкупили")
# funnel boxes
for i,(v,l) in enumerate([("297","показов"),("48","кликов"),("10","заявок")]):
 x=0.8+i*2.55
 box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(1.45),Inches(2.1),Inches(1.25)); box.fill.solid(); box.fill.fore_color.rgb=WHITE; box.line.color.rgb=LINE
 t=s.shapes.add_textbox(Inches(x),Inches(1.58),Inches(2.1),Inches(0.5)); para=t.text_frame.paragraphs[0]; para.alignment=PP_ALIGN.CENTER; r=para.add_run(); r.text=v; r.font.name=FONT; r.font.size=Pt(26); r.font.bold=True; r.font.color.rgb=RED
 t=s.shapes.add_textbox(Inches(x),Inches(2.12),Inches(2.1),Inches(0.3)); para=t.text_frame.paragraphs[0]; para.alignment=PP_ALIGN.CENTER; r=para.add_run(); r.text=l; r.font.name=FONT; r.font.size=Pt(10); r.font.color.rgb=GRAY
 if i<2:
  ar=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+2.15),Inches(1.92),Inches(0.35),Inches(0.28)); ar.fill.solid(); ar.fill.fore_color.rgb=RED; ar.line.fill.background()
bullets(s,0.55,3.05,6.0,3.5,[
 "Клик $0,36; 48 кликов за $17,28.",
 "Конверсия клика 20,8% — сильная.",
 "98% кликов с мобильных: 47 из 48.",
 "Десктоп: 1 клик, 0 заявок, $0,20 — можно не трогать.",
],size=9.8,gap=8)
text(s,6.85,3.05,5.9,0.3,"КАЧЕСТВО ТРАФИКА",9.2,True,RED)
bullets(s,6.85,3.38,5.9,3.2,[
 "Заявки — исходящие клики по номеру, не формы.",
 "6 заявок 26-го и 3 подтверждённых звонка 27-го.",
 "Нет бюджета-лимита: lostB 0% в оба дня.",
 "Ограничение — ранг/видимость, а не деньги.",
],size=9.6,gap=7)

# 5 losses
s=slide("ПОТЕРИ","Где теряем: бюджет или ранг?")
lost=[["День","IS","lostB","lostR","Верх стр.","Abs top","Потеряно показов","Потенц. клики","Потенц. заявки"],
      ["26.10","58,7%","0%","41,3%","51,3%","18,0%","~104","~18","~4,2"],
      ["27.10","61,1%","0%","38,9%","54,9%","18,1%","~95","~15","~2,5"],
      ["Итого","~60%","0%","~40%","~53%","~18%","~199","~32","~6,7"]]
table(s,0.55,1.35,12.2,2.0,lost,[1.1,0.9,0.9,0.9,1.2,1.2,1.6,1.3,1.4],font=8.8,header_font=7.6,row_h=0.38)
bullets(s,0.55,3.7,6.0,3.0,[
 "Бюджет не режет: lostB 0% оба дня.",
 "Ранг режет 41% и 39% показов.",
 "Верх страницы: 51–55%; абсолютный топ — только 18%.",
 "Потеря top-показов по рангу: 49% и 45%.",
],size=9.6,gap=8)
bullets(s,6.85,3.7,5.9,3.0,[
 "Оценка сверху: до +6–7 заявок, если выкупить весь lostR.",
 "Реалистично: половина → +3 заявки за 2 дня.",
 "Это не «добавить бюджет», это поднять ранг/QS.",
 "Снижение целевой CPA не вернёт эти показы.",
],size=9.6,gap=8)

# 6 target CPA
s=slide("ЦЕЛЕВАЯ CPA","Стоит ли снижать целевую CPA?")
kpi(s,0.55,1.35,2.8,"$16,00","текущая целевая CPA")
kpi(s,3.50,1.35,2.8,"$1,73","факт CPA",True)
kpi(s,6.45,1.35,2.8,"9,2x","запас цели к факту")
kpi(s,9.40,1.35,3.0,"40%","потеря по рангу")
bullets(s,0.55,2.75,6.0,3.8,[
 "Цель $16 — не ограничивает: факт $1,73.",
 "Снижение до $10–12 всё ещё выше факта: эффект будет почти нулевой.",
 "Снижение до $2–3 может урезать показы и не помочь рангу.",
 "Если хочешь тест — $10 на 3–5 дней, не ниже $4.",
],size=9.7,gap=8)
text(s,6.85,2.75,5.9,0.3,"ЧТО СДЕЛАТЬ ВМЕСТО",9.2,True,RED)
bullets(s,6.85,3.08,5.9,3.4,[
 "Поднять ранг: ставки/QS/объявления, а не снижать CPA.",
 "Добрать top-долю: 51–55% → 65%+.",
 "Оставить целевую CPA $16 или тест $10; не $2.",
 "Следить за CPA и объёмом 3–5 дней.",
],size=9.6,gap=7)

# 7 conclusion
s=slide("ВЫВОД","Регулировка фар: дешёвый трафик, потерянный ранг")
bullets(s,0.55,1.35,12.2,4.8,[
 "За 2 дня: 297 показов, 48 кликов, 10 заявок, $17,28, CPA $1,73.",
 "Это дешёвый и качественный трафик: CPC $0,36, CVR 20,8%, 98% мобайл.",
 "Бюджет не режет (lostB 0%), но 40% показов теряются из-за ранга.",
 "Потенциал при восстановлении половины lostR: ~+3 заявки за 2 дня.",
 "Снижать целевую CPA сейчас не нужно: цель $16 против факта $1,73; риск — потерять объём.",
 "Если тест — $10, не ниже $4; главный рычаг — ранг/QS и top-доля.",
],size=10.5,gap=10)

p.save("/Users/maksim/Desktop/Реклама/03_анализ/Регулировка_фар_26-27.10.2026.pptx")
print("saved")
